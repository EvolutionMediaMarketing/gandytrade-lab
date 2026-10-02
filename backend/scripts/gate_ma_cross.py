"""Phase 2 gate: rebuild the moving-average crossover backtest in a spreadsheet, by formulas,
and compare it trade for trade with the app's backtester.

    python scripts/gate_ma_cross.py OUT.xlsx [--bars 600]

The spreadsheet takes only the raw prices and the settings. Every average, ATR, crossing,
fill, position size, stop-loss, cost and financing charge is an ordinary formula you can
click on. LibreOffice or Excel recalculates it; the Trades sheet shows the app's figures
beside the spreadsheet's and flags any difference.
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.backtest import engine  # noqa: E402
from app.backtest.costs import Costs  # noqa: E402
from app.market.fx import Converter  # noqa: E402
from app.market.providers import sample  # noqa: E402
from app.market.symbols import get_symbol  # noqa: E402
from app.market.timeframes import get_timeframe  # noqa: E402
from app.risk.guard import RiskSettings  # noqa: E402
from app.strategies.library import STRATEGIES  # noqa: E402

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BOLD = Font(name=FONT, bold=True)
NORMAL = Font(name=FONT)
HEAD_FILL = PatternFill("solid", fgColor="DDE7F0")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")

SETTINGS = {
    "Start balance (£)": 1000.0,
    "Risk per trade": 0.01,
    "Fast average (candles)": 20,
    "Slow average (candles)": 50,
    "Stop distance (× ATR 14)": 2.0,
    "Spread (%)": 0.02,
    "Slippage (%)": 0.01,
    "Commission (£ per order)": 1.0,
    "Overnight financing (% a year)": 6.5,
    "Leverage cap (× balance)": 30,
}


def run_app(bars):
    s = engine.Settings(
        start_balance=SETTINGS["Start balance (£)"], mode="cfd", direction="long",
        risk=RiskSettings(SETTINGS["Risk per trade"] * 100, 50.0, 60.0),  # limits wide so they never trigger
        leverage=SETTINGS["Leverage cap (× balance)"],
        costs=Costs(spread_pct=SETTINGS["Spread (%)"], slippage_pct=SETTINGS["Slippage (%)"],
                    commission_gbp=SETTINGS["Commission (£ per order)"],
                    financing_pct_year=SETTINGS["Overnight financing (% a year)"]),
    )
    params = {"fast": SETTINGS["Fast average (candles)"], "slow": SETTINGS["Slow average (candles)"],
              "stop_atr": SETTINGS["Stop distance (× ATR 14)"]}
    result = engine.run(bars, STRATEGIES["ma_cross"], params, s, Converter("GBP"))
    limit_blocks = [k for k in result.skipped if "limit" in k.lower()]
    assert not result.halted and not limit_blocks, "risk limits fired; the spreadsheet doesn't model them"
    return result


def build(out: Path, n_bars: int) -> dict:
    sym = get_symbol("EUR_GBP")  # priced in pounds, so no currency conversion is involved
    bars = sample.generate(sym, get_timeframe("1d"), n_bars, now=1_790_000_000)
    result = run_app(bars)

    wb = Workbook()
    # --- Read me ---
    ws = wb.active
    ws.title = "Read me"
    lines = [
        ("Phase 2 gate: moving-average crossover, checked by hand", BOLD),
        ("", NORMAL),
        ("The 'Calc' sheet replays the strategy one candle per row using only the prices and the settings.", NORMAL),
        ("Each candle: (1) at the open, carry out orders decided at the previous close: exits first, then entries;", NORMAL),
        ("(2) during the candle, close the trade if the low reaches the stop-loss (at the open if it gapped below);", NORMAL),
        ("(3) at the close, decide what to do at the next open.", NORMAL),
        ("Buy price = open × (1 + spread/2 + slippage). Sell price = price × (1 − spread/2 − slippage).", NORMAL),
        ("Size = cash × risk ÷ (buy price − stop), capped so the position is at most 30× cash.", NORMAL),
        ("Financing = units × open price at entry × rate ÷ 365 × nights held. Commission on every order.", NORMAL),
        ("The 'Trades' sheet puts the app's results beside the spreadsheet's. The 'Summary' sheet (last tab) shows whether they all match.", NORMAL),
        ("Prices: made-up sample data for EUR/GBP (priced in pounds). The arithmetic is what's being checked.", NORMAL),
        ("Yellow cells on 'Inputs' are the settings; change one and both columns would need re-running the app to compare.", NORMAL),
    ]
    for i, (text, font) in enumerate(lines, 1):
        ws.cell(i, 1, text).font = font
    ws.column_dimensions["A"].width = 120

    # --- Inputs ---
    inp = wb.create_sheet("Inputs")
    inp["A1"], inp["B1"] = "Setting", "Value"
    inp["A1"].font = inp["B1"].font = BOLD
    names = {}
    for i, (k, v) in enumerate(SETTINGS.items(), 2):
        inp.cell(i, 1, k).font = NORMAL
        c = inp.cell(i, 2, v)
        c.font, c.fill = BLUE, INPUT_FILL
        names[k] = f"Inputs!$B${i}"
    row = len(SETTINGS) + 2
    inp.cell(row, 1, "Price adjustment per order (spread/2 + slippage)").font = NORMAL
    inp.cell(row, 2, f"={names['Spread (%)']}/200+{names['Slippage (%)']}/100").font = NORMAL
    ADJ = f"Inputs!$B${row}"
    inp.cell(row + 1, 1, "Financing per £ per night").font = NORMAL
    inp.cell(row + 1, 2, f"={names['Overnight financing (% a year)']}/100/365").font = NORMAL
    FIN = f"Inputs!$B${row + 1}"
    inp.column_dimensions["A"].width = 48
    inp.column_dimensions["B"].width = 14
    START, RISK, FAST, SLOW, KATR = (names["Start balance (£)"], names["Risk per trade"], names["Fast average (candles)"],
                                     names["Slow average (candles)"], names["Stop distance (× ATR 14)"])
    COMM, CAP = names["Commission (£ per order)"], names["Leverage cap (× balance)"]

    # --- Calc ---
    cs = wb.create_sheet("Calc")
    headers = [
        ("A", "Date"), ("B", "Open"), ("C", "High"), ("D", "Low"), ("E", "Close"),
        ("F", "Fast avg"), ("G", "Slow avg"), ("H", "True range"), ("I", "ATR 14"),
        ("J", "Crossed up"), ("K", "Crossed down"), ("L", "Stop if bought"),
        ("M", "Exit at open?"), ("N", "Sell fill at open"), ("O", "Cash change: exit at open"), ("P", "Cash after open exit"),
        ("Q", "Buy at open?"), ("R", "Buy fill"), ("S", "Units bought"), ("T", "Units held"), ("U", "Stop-loss"),
        ("V", "Stop hit?"), ("W", "Cash change: stop-loss"), ("X", "Cash at close"), ("Y", "Units at close"),
        ("Z", "Sell at next open?"), ("AA", "Buy at next open?"), ("AB", "Entry fill"), ("AC", "Entry date"),
        ("AD", "Entry open"), ("AE", "Stop for next buy"), ("AF", "Stop carried"), ("AG", "Trade no."),
        ("AH", "Trade no. before open"), ("AI", "Equity"), ("AJ", "Cash change: end of test"),
        ("AK", "Sell fill if stopped"),
    ]
    for col, name in headers:
        c = cs[f"{col}1"]
        c.value, c.font, c.fill = name, BOLD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    # Row 2: the starting state, before the first candle.
    cs["A2"] = "Start"
    cs["X2"] = f"={START}"
    for col in ("Y", "AB", "AD", "AG"):
        cs[f"{col}2"] = 0
    cs["AC2"] = 0
    cs["Z2"] = cs["AA2"] = False
    cs["AE2"] = cs["AF2"] = ""

    first, last = 3, 3 + len(bars) - 1
    for i, b in enumerate(bars):
        r, p = first + i, first + i - 1
        f = {
            "A": b.ts / 86400 + 25569,
            "B": b.open, "C": b.high, "D": b.low, "E": b.close,
            "F": f"=IF(ROW()-2>={FAST},AVERAGE(OFFSET(E{r},1-{FAST},0,{FAST},1)),\"\")",
            "G": f"=IF(ROW()-2>={SLOW},AVERAGE(OFFSET(E{r},1-{SLOW},0,{SLOW},1)),\"\")",
            "H": f"=C{r}-D{r}" if i == 0 else f"=MAX(C{r}-D{r},ABS(C{r}-E{p}),ABS(D{r}-E{p}))",
            "I": f"=IF(ROW()-2<14,\"\",IF(ROW()-2=14,AVERAGE(OFFSET(H{r},-13,0,14,1)),(I{p}*13+H{r})/14))",
            "J": f"=IF(AND(ISNUMBER(F{r}),ISNUMBER(G{r}),ISNUMBER(F{p}),ISNUMBER(G{p})),AND(F{r}>G{r},F{p}<=G{p}),FALSE)",
            "K": f"=IF(AND(ISNUMBER(F{r}),ISNUMBER(G{r}),ISNUMBER(F{p}),ISNUMBER(G{p})),AND(F{r}<G{r},F{p}>=G{p}),FALSE)",
            "L": f"=IF(ISNUMBER(I{r}),E{r}-{KATR}*I{r},\"\")",
            "M": f"=AND(Y{p}>0,Z{p})",
            "N": f"=B{r}*(1-{ADJ})",
            "O": f"=IF(M{r},Y{p}*(N{r}-AB{p})-{COMM}-Y{p}*AD{p}*{FIN}*(A{r}-AC{p}),0)",
            "P": f"=X{p}+O{r}",
            "Q": f"=AND(IF(M{r},0,Y{p})=0,AA{p},IF(ISNUMBER(AE{p}),B{r}>AE{p},FALSE))",
            "R": f"=B{r}*(1+{ADJ})",
            "S": f"=IF(Q{r},MIN(P{r}*{RISK}/(R{r}-AE{p}),P{r}*{CAP}/R{r}),0)",
            "T": f"=IF(Q{r},S{r},IF(M{r},0,Y{p}))",
            "U": f"=IF(Q{r},AE{p},AF{p})",
            "V": f"=AND(T{r}>0,D{r}<=U{r})",
            "AK": f"=IF(AND(B{r}<=U{r},NOT(Q{r})),B{r},U{r})*(1-{ADJ})",
            "W": f"=IF(V{r},T{r}*(AK{r}-AB{r})-{COMM}-T{r}*AD{r}*{FIN}*(A{r}-AC{r}),0)",
            "X": f"=P{r}-IF(Q{r},{COMM},0)+W{r}",
            "Y": f"=IF(V{r},0,T{r})",
            "Z": f"=AND(Y{r}>0,K{r})",
            "AA": f"=AND(Y{r}=0,J{r},ISNUMBER(L{r}))",
            "AB": f"=IF(Q{r},R{r},AB{p})",
            "AC": f"=IF(Q{r},A{r},AC{p})",
            "AD": f"=IF(Q{r},B{r},AD{p})",
            "AE": f"=IF(AA{r},L{r},\"\")",
            "AF": f"=U{r}",
            "AG": f"=AG{p}+IF(Q{r},1,0)",
            "AH": f"=AG{p}",
            "AI": f"=X{r}+Y{r}*(E{r}-AB{r})-Y{r}*AD{r}*{FIN}*(A{r}-AC{r})",
            "AJ": (f"=IF(Y{r}>0,Y{r}*(E{r}*(1-{ADJ})-AB{r})-{COMM}-Y{r}*AD{r}*{FIN}*(A{r}-AC{r}),0)" if r == last else 0),
        }
        for col, val in f.items():
            c = cs[f"{col}{r}"]
            c.value = val
            c.font = BLUE if col in "ABCDE" and len(col) == 1 else NORMAL
        cs[f"A{r}"].number_format = "dd mmm yyyy"
        for col in ("B", "C", "D", "E", "F", "G", "H", "I", "L", "N", "R", "U", "AB", "AD", "AE", "AF", "AK"):
            cs[f"{col}{r}"].number_format = "0.000000"
        for col in ("O", "P", "W", "X", "AI", "AJ"):
            cs[f"{col}{r}"].number_format = "£#,##0.00;-£#,##0.00;-"
        for col in ("S", "T", "Y"):
            cs[f"{col}{r}"].number_format = "#,##0.00"
        cs[f"AC{r}"].number_format = "dd mmm yyyy"
    for idx in range(1, 38):
        cs.column_dimensions[get_column_letter(idx)].width = 13
    cs.freeze_panes = "B2"

    # --- Trades ---
    ts = wb.create_sheet("Trades")
    th = ["Trade", "Entry date (app)", "Entry date (sheet)", "Buy fill (app)", "Buy fill (sheet)", "Units (app)",
          "Units (sheet)", "Exit (app)", "Profit £ (app)", "Profit £ (sheet)", "Difference £", "Match?"]
    for j, h in enumerate(th, 1):
        c = ts.cell(1, j, h)
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True)
    rng = f"Calc!$AG${first}:$AG${last}"
    for k, t in enumerate(result.trades, 1):
        r = k + 1
        row_of = f"MATCH({k},{rng},0)"
        ts.cell(r, 1, k)
        ts.cell(r, 2, t.entry_ts / 86400 + 25569).number_format = "dd mmm yyyy"
        ts.cell(r, 3, f"=INDEX(Calc!$A${first}:$A${last},{row_of})").number_format = "dd mmm yyyy"
        ts.cell(r, 4, t.entry_price).number_format = "0.000000"
        ts.cell(r, 5, f"=INDEX(Calc!$R${first}:$R${last},{row_of})").number_format = "0.000000"
        ts.cell(r, 6, t.units).number_format = "#,##0.0000"
        ts.cell(r, 7, f"=INDEX(Calc!$S${first}:$S${last},{row_of})").number_format = "#,##0.0000"
        ts.cell(r, 8, t.exit_reason)
        ts.cell(r, 9, round(t.pnl_gbp, 6)).number_format = "£#,##0.00;-£#,##0.00"
        ts.cell(r, 10, (f"=SUMIF(Calc!$AH${first}:$AH${last},{k},Calc!$O${first}:$O${last})"
                        f"+SUMIF({rng},{k},Calc!$W${first}:$W${last})"
                        f"+SUMIF({rng},{k},Calc!$AJ${first}:$AJ${last})-{COMM}")).number_format = "£#,##0.00;-£#,##0.00"
        ts.cell(r, 11, f"=J{r}-I{r}").number_format = "£0.000000;-£0.000000"
        ts.cell(r, 12, f"=AND(ABS(K{r})<0.005,C{r}=B{r},ABS(E{r}-D{r})<0.0000001,ABS(G{r}-F{r})<0.0001)")
        for j in (2, 4, 6, 8, 9):
            ts.cell(r, j).font = BLUE
    for j, w in enumerate([7, 14, 14, 12, 12, 12, 12, 40, 13, 13, 13, 9], 1):
        ts.column_dimensions[get_column_letter(j)].width = w
    ts.freeze_panes = "A2"
    n = len(result.trades)

    # --- Summary ---
    sm = wb.create_sheet("Summary")
    rows = [
        ("", "App", "Spreadsheet", "Match?"),
        ("Number of trades", n, f"=MAX(Calc!$AG${first}:$AG${last})", "=B2=C2"),
        ("Final balance (£)", round(result.equity[-1][1], 6), f"=Calc!X{last}+Calc!AJ{last}", "=ABS(B3-C3)<0.01"),
        ("Trades matching", n, f"=COUNTIF(Trades!L2:L{n + 1},TRUE)", "=B4=C4"),
        ("GATE", "", "", "=AND(D2,D3,D4)"),
    ]
    for i, vals in enumerate(rows, 1):
        for j, v in enumerate(vals, 1):
            c = sm.cell(i, j, v)
            c.font = BOLD if i == 1 or j == 1 else (BLUE if j == 2 else NORMAL)
    sm["B3"].number_format = sm["C3"].number_format = "£#,##0.00"
    sm.column_dimensions["A"].width = 24
    for col in "BCD":
        sm.column_dimensions[col].width = 16
    wb.save(out)
    return {"trades": n, "final": result.equity[-1][1], "bars": len(bars)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--bars", type=int, default=600)
    a = ap.parse_args()
    print(build(Path(a.out), a.bars))
